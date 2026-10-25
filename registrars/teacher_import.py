"""The registrar's teacher list for creating many teacher accounts at once (Teachers > Import teachers).

Three kinds of file work:
- DepEd School Form 7 (SF7, School Personnel Assignment List and Basic Profile), the form every school prepares at
  the start of the school year. ARMS reads its personnel table (employee number, name, position, and major as the
  teaching area) and keeps the teaching positions only. SF7 has no email column, so each teacher gets
  firstname.lastname@<school domain>, as a blank email does in the other files.
- The ARMS teacher template (.xlsx, downloadable on the page): First name, Last name, Email, Employee number,
  Position, Teaching area.
- A CSV with the same columns (the earlier format; the old first_name, last_name, ... headers still work).
Nothing is created here: the rows go to the review page first (registrars.views.teacher_import).
"""
import csv
import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

MAX_ROWS = 500
MAX_CSV_BYTES = 1024 * 1024
MAX_EXCEL_BYTES = 5 * 1024 * 1024
FIELDS = ('first_name', 'last_name', 'email', 'employee_number', 'position', 'teaching_area')
# header words, lower case and without "(optional)", that name each field
ALIASES = {
    'first_name': ('first_name', 'first name', 'firstname', 'given name', 'given names'),
    'last_name': ('last_name', 'last name', 'lastname', 'surname', 'family name'),
    'email': ('email', 'email address', 'e-mail', 'e-mail address', 'deped email'),
    'employee_number': ('employee_number', 'employee number', 'employee no', 'employee no.', 'employee id', 'employee id number'),
    'position': ('position', 'designation', 'position/designation', 'position title', 'plantilla position'),
    'teaching_area': ('teaching_area', 'teaching area', 'department', 'subject area', 'major', 'specialization', 'learning area'),
}
TEACHING = re.compile(r'teacher|instructor|professor|lecturer|faculty', re.I)
TEMPLATE_HEADERS = ('First name', 'Last name', 'Email (optional)', 'Employee number (optional)', 'Position', 'Teaching area')
POSITIONS = ('Teacher I', 'Teacher II', 'Teacher III', 'Teacher IV', 'Teacher V', 'Teacher VI', 'Teacher VII',
             'Master Teacher I', 'Master Teacher II', 'Master Teacher III', 'Master Teacher IV',
             'Head Teacher I', 'Head Teacher II', 'Head Teacher III')


class TeacherFileError(Exception):
    """The file cannot be used (wrong kind, unreadable, too big or empty)."""


def _key(header):
    text = re.sub(r'\(.*?\)', '', str(header or '')).strip().lower()
    text = re.sub(r'\s+', ' ', text.replace('_', ' ')).strip(' :*')
    for field, names in ALIASES.items():
        if text in names or text.replace(' ', '_') in names:
            return field
    return None


def _cell(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))                 # employee numbers typed as numbers
    return ' '.join(str(value).split())


def _named(name):
    """'DELA CRUZ, MARIA CLARA S.' -> ('Maria Clara', 'Dela Cruz'); middle names and initials are dropped."""
    text = ' '.join(name.replace('\n', ' ').split())
    if ',' in text:
        last, rest = (part.strip() for part in text.split(',', 1))
        words = [w for w in rest.split() if len(w.strip('.')) > 1]
        given = ' '.join(words[:-1] if len(words) > 2 else words) if words else rest
    else:
        words = text.split()
        if len(words) < 2:
            return _proper(text), ''
        given = ' '.join(w for w in words[:-1] if len(w.strip('.')) > 1) or words[0]
        last = words[-1]
    return _proper(given), _proper(last)


def _proper(text):
    """'DELA CRUZ' -> 'Dela Cruz', keeping suffixes such as III in capitals."""
    return ' '.join(w.upper() if w.upper().strip('.') in ('II', 'III', 'IV', 'V') else w.capitalize() for w in text.split())


def _from_rows(rows):
    """Rows of a template or CSV sheet: the first row that names first and last name is the header."""
    for i, row in enumerate(rows[:12]):
        keys = [_key(v) for v in row]
        if 'first_name' in keys and 'last_name' in keys:
            out = []
            for values in rows[i + 1:]:
                item = {key: _cell(values[j]) for j, key in enumerate(keys) if key and j < len(values)}
                if any(item.values()):
                    out.append({field: item.get(field, '') for field in FIELDS})
            return out
    return None


def _sf7(rows):
    """(teachers, skipped) from the personnel table of an SF7 sheet, or None when the sheet is not SF7."""
    for i, row in enumerate(rows):
        texts = [str(v or '').upper() for v in row]
        name_col = next((j for j, t in enumerate(texts) if 'NAME OF SCHOOL PERSONNEL' in t), None)
        if name_col is None:
            continue
        below = [str(v or '').upper() for v in (rows[i + 1] if i + 1 < len(rows) else [])]
        # a header is found by how it starts: the name column's own header says "(Arrange by Position, Descending)"
        find = lambda word, cells: next((j for j, t in enumerate(cells) if t.strip().startswith(word)), None)
        employee = find('EMPLOYEE NO', texts)
        position = find('POSITION', texts)
        major = find('MAJOR', below) if find('MAJOR', below) is not None else find('MAJOR', texts)
        subject = find('SUBJECT', texts)
        teachers, skipped = [], []
        for values in rows[i + 2:]:
            get = lambda j: _cell(values[j]) if j is not None and j < len(values) else ''
            name, title = get(name_col), get(position)
            if not name or not re.search(r'[A-Za-z]{2}', name) or not title:
                continue                               # the other rows of a person's block, and the footer
            if re.search(r'prepared|certified|school head|attested|name of', name, re.I):
                continue
            if not TEACHING.search(title):
                skipped.append({'name': name, 'position': title, 'reason': 'Not a teaching position'})
                continue
            first, last = _named(name)
            area = get(major) or re.sub(r'\s*\(.*?\)\s*', ' ', get(subject)).strip()
            teachers.append({'first_name': first, 'last_name': last, 'email': '', 'employee_number': get(employee),
                             'position': title, 'teaching_area': area[:100]})
        return teachers, skipped
    return None


def _excel_rows(upload, name):
    data = upload.read()
    if name.endswith('.xls'):
        import xlrd
        try:
            book = xlrd.open_workbook(file_contents=data)
        except Exception as exc:
            raise TeacherFileError('That Excel file could not be opened. Save it again as .xlsx and try once more.') from exc
        return [[[sheet.cell_value(r, c) for c in range(sheet.ncols)] for r in range(sheet.nrows)] for sheet in book.sheets()]
    from openpyxl import load_workbook
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise TeacherFileError('That Excel file could not be opened. Check that it is an .xlsx file.') from exc
    sheets = [[list(r) for r in ws.iter_rows(values_only=True, max_row=2000, max_col=40)] for ws in book.worksheets]
    book.close()
    return sheets


def read(upload):
    """(rows, skipped, kind) from an uploaded file: kind is 'SF7', 'template' or 'CSV'. Raises TeacherFileError."""
    if not upload:
        raise TeacherFileError('Choose the teacher list first: SF7, the ARMS template (.xlsx) or a CSV.')
    name = upload.name.lower()
    if name.endswith('.csv'):
        if upload.size > MAX_CSV_BYTES:
            raise TeacherFileError('That CSV is larger than 1 MB. A teacher list is much smaller; check the file.')
        try:
            text = upload.read().decode('utf-8-sig')
        except UnicodeDecodeError as exc:
            raise TeacherFileError('That file is not a readable CSV. Save it as "CSV UTF-8" and try again.') from exc
        rows, skipped, kind = _from_rows(list(csv.reader(io.StringIO(text)))), [], 'CSV'
        if rows is None:
            raise TeacherFileError('The CSV needs First name and Last name columns. Download the template to see the format.')
    elif name.endswith(('.xlsx', '.xlsm', '.xls')):
        if upload.size > MAX_EXCEL_BYTES:
            raise TeacherFileError('That Excel file is larger than 5 MB. A teacher list or SF7 is much smaller; check the file.')
        rows = skipped = kind = None
        for sheet in _excel_rows(upload, name):
            found = _sf7(sheet)
            if found is not None:
                (rows, skipped), kind = found, 'SF7'
                break
            found = _from_rows(sheet)
            if found is not None:
                rows, skipped, kind = found, [], 'template'
                break
        if rows is None:
            raise TeacherFileError('This file is neither an SF7 nor the ARMS teacher template: no sheet has the personnel '
                                   'table of SF7 or First name and Last name columns.')
    else:
        raise TeacherFileError('Upload SF7 or the ARMS template (.xlsx or .xls), or a CSV (.csv).')
    if not rows:
        raise TeacherFileError('No teachers were found in the file.' + (' Its personnel are all in non-teaching positions.' if skipped else ''))
    if len(rows) > MAX_ROWS:
        raise TeacherFileError(f'Please import at most {MAX_ROWS} teachers at a time.')
    return rows, skipped, kind


def template_xlsx(school_name, email_domain):
    """The ARMS teacher template: one sheet to fill in, one with the instructions."""
    book = Workbook()
    sheet = book.active
    sheet.title = 'Teachers'
    sheet.append(TEMPLATE_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='123499')
        cell.alignment = Alignment(vertical='center')
    for column, width in zip('ABCDEF', (22, 20, 34, 30, 20, 24)):
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 22
    sheet.freeze_panes = 'A2'
    positions = DataValidation(type='list', formula1='"' + ','.join(POSITIONS) + '"', allow_blank=True,
                               errorStyle='warning', error='Choose a DepEd teaching position, or type another one.')
    sheet.add_data_validation(positions)
    positions.add('E2:E501')
    notes = book.create_sheet('How to fill it')
    notes.column_dimensions['A'].width = 110
    lines = [
        f'ARMS teacher template for {school_name}',
        '',
        'One row per teacher, starting on row 2 of the Teachers sheet. Do not rename or move the column headers.',
        'First name and Last name are required. Leave out middle names and initials.',
        (f'Email: leave it empty to give the teacher firstname.lastname@{email_domain}. The email is also the username.'
         if email_domain else 'Email: required, because the school has no email domain in ARMS. The email is also the username.'),
        'Employee number: optional, but no two teachers may share one.',
        'Position: choose from the list (Teacher I to Head Teacher III) or type another; empty means Teacher.',
        'Teaching area: the learning area the teacher handles most, for example Mathematics.',
        '',
        'Upload it on Teachers > Import teachers. Nothing is created until you confirm the review page.',
        'Each new teacher gets a temporary password, shown once, and sets their own at first sign-in.',
        'You can also upload the school\'s SF7 (School Personnel Assignment List) instead: ARMS reads its teaching staff.',
    ]
    for line in lines:
        notes.append([line])
    notes['A1'].font = Font(bold=True, size=13, color='123499')
    data = io.BytesIO()
    book.save(data)
    return data.getvalue()
