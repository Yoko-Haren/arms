"""Ready-made DepEd e-class records for one class.

The official template (grades/ecr_templates) with the class filled in: school, grade and section,
teacher, and the learners in DepEd order (male, then female, "LAST, FIRST M."), each with their
LRN in a hidden column of INPUT DATA so an upload can be tied to the exact learner.

Grades 7 to 10: one template per learning area (MAPEH one per quarter). Any sample names and
scores are cleared. For trimestral or semestral schools the period sheets are relabelled, unused
ones hidden, and the summary's final grade averages only the school's periods.

Senior High: the Strengthened SHS e-class record (SSHS.xlsx). It leans on Excel 365 features
(dynamic-array helper formulas, dropdowns that read other sheets), so it is filled cell by cell in
its XML (core.xlsx_patch) and every other part of DepEd's file stays exactly as it is. The subject
type that sets the weights comes from the template's own subject list, else from the ARMS subject.

Grades 7 to 10 of a school year graded in three terms: DepEd's three-term record of DepEd Order No. 015,
s. 2026 (TERM_*.xlsx): one for the core learning areas, one for GMRC and Values Education, one for
MAPEH (Music and Arts, PE and Health on two sheets a term) and one for EPP/TLE, filled in its XML the
same way, so its protected sheets and array formulas stay as DepEd made them.
"""
import functools
import io
import re
from pathlib import Path

from openpyxl import load_workbook

from core import periods as P
from core import school_header, xlsx_patch

from . import ecr

TEMPLATE_DIR = Path(__file__).resolve().parent / 'ecr_templates'
FILES = {'MATH': 'MATH.xlsx', 'SCIENCE': 'SCIENCE.xlsx', 'ENGLISH': 'ENGLISH.xlsx', 'FILIPINO': 'FILIPINO.xlsx',
         'AP': 'AP.xlsx', 'ESP': 'ESP.xlsx', 'TLE': 'TLE.xlsx'}
TERM_FILES = {'MATH': 'TERM_CORE.xlsx', 'SCIENCE': 'TERM_CORE.xlsx', 'ENGLISH': 'TERM_CORE.xlsx', 'FILIPINO': 'TERM_CORE.xlsx',
              'AP': 'TERM_CORE.xlsx', 'ESP': 'TERM_GMRC.xlsx', 'TLE': 'TERM_TLE.xlsx', 'MAPEH': 'TERM_MAPEH.xlsx'}
SSHS_FILE = 'SSHS.xlsx'
SSHS_OTHER = 'Other Elective and Special Curricular Program'
_TECHPRO_STRANDS = {'ICT', 'HE', 'IA', 'AFA', 'TVL'}
MALE_ROWS = range(12, 62)            # INPUT DATA and period sheets; the summary is one row lower
FEMALE_ROWS = range(63, 113)
SUMMARY_PERIOD_COLS = ('F', 'J', 'N', 'R')
ORDINAL_SHORT = {1: '1st', 2: '2nd', 3: '3rd', 4: '4th'}


class TemplateUnavailable(Exception):
    pass


def _is_mapeh(assignment):
    s = assignment.subject
    family = s.subject_family.family_name if s.subject_family_id else ''
    return any(re.search(r'\bMAPEH\b', (t or '').upper()) for t in (s.subject_name, family, s.subject_code))


def template_for(assignment, period_number):
    if assignment.section.grade_level.is_senior_high:
        return TEMPLATE_DIR / SSHS_FILE, False
    if _is_mapeh(assignment):
        return TEMPLATE_DIR / f'MAPEH_{min(max(period_number, 1), 4)}.xlsx', True
    s = assignment.subject
    family = s.subject_family.family_name if s.subject_family_id else ''
    key = ecr.subject_key(s.subject_name) or ecr.subject_key(family)
    if key not in FILES:
        raise TemplateUnavailable(f'There is no DepEd e-class record template for {s.subject_name}. '
                                  'Download the class list (CSV) instead.')
    return TEMPLATE_DIR / FILES[key], False


@functools.lru_cache(maxsize=1)
def sshs_subjects():
    """The SSHS template's own subject list: {clean name: (subject type, official name)}."""
    wb = load_workbook(TEMPLATE_DIR / SSHS_FILE, read_only=True, data_only=True)
    try:
        out = {}
        for kind, _, name, *_ in wb['Helper(IMPORTANT!)'].iter_rows(min_row=2, max_row=205, min_col=113, max_col=118,
                                                                    values_only=True):       # DI:DN
            if kind and name and not str(name).startswith(SSHS_OTHER):
                out.setdefault(ecr._clean(name), (str(kind).strip(), str(name).strip()))
        return out
    finally:
        wb.close()


def sshs_subject_type(subject):
    """The SSHS subject type (and so the weights) of an ARMS subject."""
    name = ecr._clean(subject.subject_name)
    listed = sshs_subjects().get(name)
    if listed and listed[0] in ecr.SSHS_TYPES:
        return listed[0]
    if 'WORK IMMERSION' in name:
        return 'TechPro Elective (Work Immersion)'
    if 'FIELD EXPOSURE' in name or 'FIELD EXPERIENCE' in name:
        return 'Academic Elective (Field Experience / Exposure, and Sports and Arts)'
    if subject.subject_category == 'Core':
        return 'Core Subject (All Tracks)'
    strand = subject.strand if subject.strand_id else None
    track = strand.track.track_code.upper() if strand and strand.track_id else ''
    if track in ('TVL', 'TECHPRO') or (strand and strand.strand_code.upper() in _TECHPRO_STRANDS):
        return 'TechPro Elective (All Other Electives)'
    return 'Academic Elective (All Other Electives)'


def learner_name(p):
    """DELA CRUZ, JUAN MIGUEL S. (the DepEd way of writing a learner in the class record)."""
    ext = f' {p["ext"]}' if p.get('ext') else ''
    middle = f' {p["middle"].strip()[0]}.' if p.get('middle', '').strip() else ''
    return f'{p["last"]}{ext}, {p["first"]}{middle}'.upper()


def _grid_of(ws):
    rows = [[None]]
    for row in ws.iter_rows(values_only=True):
        rows.append([None] + list(row))
    return rows


def _period_sheets(wb):
    """The grading-period sheets (in period order) with their score columns."""
    found = []
    for ws in wb.worksheets:
        if 'INPUT' in ws.title.upper() or 'SUMMARY' in ws.title.upper() or 'DO NOT' in ws.title.upper():
            continue
        if 'HIGHEST POSSIBLE' not in str(ws['B10'].value or '').upper():
            continue
        layout = ecr._layout(_grid_of(ws))
        number = ecr._period_number(ws.title, str(ws['A7'].value or ''))
        if layout and number:
            found.append((number, ws, layout))
    return sorted(found, key=lambda x: (x[0], x[1].title))


def build(assignment, roster, teacher, period=None):
    """(bytes, file name) of the class's e-class record."""
    from academics.models import Quarter
    school, sy = assignment.section.school, assignment.school_year
    periods = list(Quarter.objects.filter(school_year=sy).order_by('quarter_number'))
    own = P.for_class(assignment, periods)
    count = len(own) or 4
    prefix = P.prefix(school, periods)
    number = period.quarter_number if period else (own[0].quarter_number if own else 1)
    path, mapeh = template_for(assignment, number)
    if not path.exists():
        raise TemplateUnavailable('The e-class record template is missing on the server. Ask the administrator to restore grades/ecr_templates.')
    males = [p for p in roster if p['sex'] == 'M']
    females = [p for p in roster if p['sex'] != 'M']
    if len(males) > 50 or len(females) > 50:
        raise TemplateUnavailable('The e-class record has room for 50 male and 50 female learners; this class has more. Use the class list (CSV).')
    if assignment.section.grade_level.is_senior_high:
        return _build_sshs(assignment, males, females, teacher, path)
    if prefix == 'T':
        return _build_term(assignment, males, females, teacher)

    wb = load_workbook(path)
    inp = next(ws for ws in wb.worksheets if 'INPUT' in ws.title.upper())
    teacher_name = (teacher.get_full_name() or teacher.username).upper()
    inp['G5'] = school.school_name.upper()
    inp['X5'] = school.school_id
    inp['AG5'] = sy.year_label
    inp['K7'] = f'{assignment.section.grade_level.grade_name} - {assignment.section.section_name}'.upper()
    inp['S7'] = teacher_name
    for r in list(MALE_ROWS) + list(FEMALE_ROWS):
        inp.cell(row=r, column=2).value = None
        inp[f'{ecr.LRN_COLUMN}{r}'] = None
    for rows, people in ((MALE_ROWS, males), (FEMALE_ROWS, females)):
        for r, p in zip(rows, people):
            inp.cell(row=r, column=2).value = learner_name(p)
            cell = inp[f'{ecr.LRN_COLUMN}{r}']
            cell.value = str(p['lrn'] or '')
            cell.number_format = '@'
    inp[f'{ecr.LRN_COLUMN}11'] = 'LRN (do not edit)'
    inp.column_dimensions[ecr.LRN_COLUMN].hidden = True

    sheets = _period_sheets(wb)
    for _, ws, layout in sheets:                         # a clean record: no sample scores
        for part in ('ww', 'pt', 'qa'):
            for col in layout[part]['items']:
                ws.cell(row=10, column=col).value = None
                for r in list(MALE_ROWS) + list(FEMALE_ROWS):
                    ws.cell(row=r, column=col).value = None

    word = P.PERIOD_WORDS.get(prefix, P.PERIOD_WORDS['Q'])[0]
    if prefix != 'Q' or count != 4:
        _relabel(wb, sheets, mapeh, prefix, word, count, number)

    buf = io.BytesIO()
    wb.save(buf)
    code = re.sub(r'[^A-Za-z0-9]+', '-', assignment.subject.subject_code).strip('-')
    section = re.sub(r'[^A-Za-z0-9]+', '-', assignment.section.section_name).strip('-')
    suffix = f'_{prefix}{number}' if mapeh else ''
    return buf.getvalue(), f'ECR_{code}_{section}{suffix}_SY{sy.year_label}.xlsx'


def _build_sshs(assignment, males, females, teacher, path):
    """The SSHS record with the class in INPUT DATA; every other part of the template untouched."""
    section, subject, sy = assignment.section, assignment.subject, assignment.school_year
    school = section.school
    place = school_header.get(school)
    kind = sshs_subject_type(subject)
    listed = sshs_subjects().get(ecr._clean(subject.subject_name))
    cells = {
        'F10': school_header.region_label(place['region']) or None, 'F11': place['division'] or None,
        'F13': school.school_id, 'F14': school.school_name.upper(), 'F16': sy.year_label,
        'F22': (teacher.get_full_name() or teacher.username).upper(),
        'F24': section.grade_level.grade_number, 'F25': section.section_name.upper(), 'F26': kind,
        # a subject of DepEd's list goes in SUBJECT; any other in "Other Elective & Special Curricular Program / Subject"
        'F28': listed[1] if listed else (None if kind == 'Core Subject (All Tracks)' else SSHS_OTHER),
        'F30': None if listed else subject.subject_name,
    }
    male_lrn, female_lrn = ecr.SSHS_LRN_COLUMNS
    cells[f'{male_lrn}{ecr.SSHS_FIRST_ROW - 1}'] = 'LRN (male, do not edit)'
    cells[f'{female_lrn}{ecr.SSHS_FIRST_ROW - 1}'] = 'LRN (female, do not edit)'
    for name_col, lrn_col, people in (('L', male_lrn, males), ('O', female_lrn, females)):
        for i in range(50):
            r = ecr.SSHS_FIRST_ROW + i
            p = people[i] if i < len(people) else None
            cells[f'{name_col}{r}'] = learner_name(p) if p else None
            cells[f'{lrn_col}{r}'] = str(p['lrn'] or '') if p else None
    content = xlsx_patch.patch(path.read_bytes(), {'INPUT DATA': cells})
    code = re.sub(r'[^A-Za-z0-9]+', '-', subject.subject_code).strip('-')
    name = re.sub(r'[^A-Za-z0-9]+', '-', section.section_name).strip('-')
    return content, f'ECR_{code}_{name}_SY{sy.year_label}.xlsx'


def _build_term(assignment, males, females, teacher):
    """DepEd's three-term record (DepEd Order No. 015, s. 2026) with the class in INPUT DATA; the rest untouched."""
    section, subject, sy = assignment.section, assignment.subject, assignment.school_year
    school = section.school
    family = subject.subject_family.family_name if subject.subject_family_id else ''
    key = 'MAPEH' if _is_mapeh(assignment) else (ecr.subject_key(subject.subject_name) or ecr.subject_key(family))
    if key not in TERM_FILES:
        raise TemplateUnavailable(f'There is no DepEd three-term e-class record for {subject.subject_name}. '
                                  'Download the class list (CSV) instead.')
    path = TEMPLATE_DIR / TERM_FILES[key]
    if not path.exists():
        raise TemplateUnavailable('The three-term e-class record template is missing on the server. Ask the administrator to '
                                  'restore grades/ecr_templates.')
    from accounts.models import UserProfile
    head = UserProfile.objects.filter(school=school, role='schoolhead', is_active=True, user__is_active=True).select_related('user').first()
    place = school_header.get(school)
    cells = {
        'E10': school_header.region_label(place['region']) or None, 'E11': place['division'] or None,
        'E13': school.school_id, 'E14': school.school_name.upper(), 'E15': sy.year_label,
        'E16': (head.user.get_full_name() or head.user.username).upper() if head else None,
        'E23': (teacher.get_full_name() or teacher.username).upper(), 'E24': subject.subject_name,
        'E25': section.grade_level.grade_name, 'E26': section.section_name.upper(),
    }
    male_lrn, female_lrn = ecr.TERM_LRN_COLUMNS
    cells[f'{male_lrn}{ecr.TERM_FIRST_ROW - 1}'] = 'LRN (male, do not edit)'
    cells[f'{female_lrn}{ecr.TERM_FIRST_ROW - 1}'] = 'LRN (female, do not edit)'
    for name_col, lrn_col, people in (('K', male_lrn, males), ('N', female_lrn, females)):
        for i in range(50):
            r = ecr.TERM_FIRST_ROW + i
            p = people[i] if i < len(people) else None
            cells[f'{name_col}{r}'] = learner_name(p) if p else None
            cells[f'{lrn_col}{r}'] = str(p['lrn'] or '') if p else None
    content = xlsx_patch.patch(path.read_bytes(), {'INPUT DATA': cells})
    code = re.sub(r'[^A-Za-z0-9]+', '-', subject.subject_code).strip('-')
    name = re.sub(r'[^A-Za-z0-9]+', '-', section.section_name).strip('-')
    return content, f'ECR_{code}_{name}_3Term_SY{sy.year_label}.xlsx'


def _relabel(wb, sheets, mapeh, prefix, word, count, number):
    """Name the periods the school's way and keep the summary to the school's number of periods."""
    upper = word.upper()
    summary = next((ws for ws in wb.worksheets if 'SUMMARY' in ws.title.upper()), None)
    if mapeh:                                            # one period per file: relabel it
        for _, ws, _ in sheets:
            ws['A7'] = f'{P.ORDINALS[number].upper()} {upper}'
        if summary:
            for col in SUMMARY_PERIOD_COLS:
                summary[f'{col}11'] = f'{ORDINAL_SHORT[number]} {word}'
            summary['V11'] = f'{ORDINAL_SHORT[number].upper()} {upper} GRADE'
        return
    names = {}
    for n, ws, _ in sheets:
        if n > count:
            ws.sheet_state = 'hidden'
            continue
        ws['A7'] = f'{P.ORDINALS[n].upper()} {upper}'
        new = re.sub(r'Q(\d)$', rf'{prefix}\1', ws.title)
        if new != ws.title:
            ws.title = new
        names[n] = ws.title
    if not summary:
        return
    used = SUMMARY_PERIOD_COLS[:count]
    for n, col in enumerate(SUMMARY_PERIOD_COLS, start=1):
        summary[f'{col}11'] = f'{ORDINAL_SHORT[n]} {word}' if n <= count else None
        summary.column_dimensions[col].hidden = n > count
    for r in list(range(13, 63)) + list(range(64, 114)):
        source = r - 1
        for n, col in enumerate(SUMMARY_PERIOD_COLS, start=1):
            summary[f'{col}{r}'] = f"='{names[n]}'!AJ{source}" if n <= count and n in names else None
        cells = ','.join(f'{c}{r}' for c in used)
        blanks = ','.join(f'{c}{r}=""' for c in used)
        fixed_blanks = ','.join(f'${c}{r}=""' for c in used)
        summary[f'V{r}'] = f'=IF(OR({blanks}),"",IF(ISERROR(ROUND(AVERAGE({cells}),0)),"",ROUND(AVERAGE({cells}),0)))'
        summary[f'Z{r}'] = f'=IF(OR({fixed_blanks},$V{r}=""),"",IF($V{r}>=75,"PASSED","FAILED"))'
