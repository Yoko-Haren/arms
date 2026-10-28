"""DepEd E-Class Record reader for grade uploads (DepEd Orders No. 8, s. 2015 and No. 015, s. 2026).

The official e-class record keeps one sheet per grading period. Each sheet holds every learner's
written work, performance task and quarterly assessment scores, the highest possible scores
(row 10; row 11 in the Senior High record) and the component weights. Three templates are read:
the Grades 7 to 10 record of four quarters (one per learning area, MAPEH with a sheet per component),
the Strengthened Senior High School (SSHS) record, whose weights follow the subject type chosen in
INPUT DATA, and the three-term record of DepEd Order No. 015, s. 2026 (written/oral works, product/
performance tasks, and examinations made of two summative tests and a term examination; MAPEH and
TLE with two component sheets per term). This module reads such a workbook, recomputes each grade from
the raw scores exactly as the template's formulas do (Excel rounding, then the transmutation
table), checks the file against the class it is uploaded for, and matches every row to exactly
one learner of that class. Nothing is saved here: the upload view shows the result, the teacher
resolves anything uncertain, and only then are grades written.
"""
import io
import re
import unicodedata
import zipfile
from decimal import ROUND_HALF_UP, Decimal
from difflib import SequenceMatcher

from openpyxl import load_workbook

MAX_UPLOAD_BYTES = 15 * 1024 * 1024
MAX_UNZIPPED_BYTES = 150 * 1024 * 1024
LRN_COLUMN = 'AK'                 # hidden column of INPUT DATA where ARMS-generated files keep each learner's LRN
PASSING = 75

# 'DO NOT DELETE'!G2:J42 of the template: lowest initial grade of each band -> quarterly grade
TRANSMUTATION = tuple((Decimal(str(low)), grade) for low, grade in (
    (0, 60), (4, 61), (8, 62), (12, 63), (16, 64), (20, 65), (24, 66), (28, 67), (32, 68), (36, 69),
    (40, 70), (44, 71), (48, 72), (52, 73), (56, 74), (60, 75), (61.6, 76), (63.2, 77), (64.8, 78),
    (66.4, 79), (68, 80), (69.6, 81), (71.2, 82), (72.8, 83), (74.4, 84), (76, 85), (77.6, 86),
    (79.2, 87), (80.8, 88), (82.4, 89), (84, 90), (85.6, 91), (87.2, 92), (88.8, 93), (90.4, 94),
    (92, 95), (93.6, 96), (95.2, 97), (96.8, 98), (98.4, 99), (100, 100)))

# Component weights for Grades 1 to 10 (written work, performance tasks, quarterly assessment)
AREA_WEIGHTS = {
    'languages': ('Languages, Araling Panlipunan and EsP', (Decimal('0.3'), Decimal('0.5'), Decimal('0.2'))),
    'math_science': ('Mathematics and Science', (Decimal('0.4'), Decimal('0.4'), Decimal('0.2'))),
    'mapeh_tle': ('MAPEH and EPP/TLE', (Decimal('0.2'), Decimal('0.6'), Decimal('0.2'))),
}
_AREA_WORDS = (       # checked in order; EPP and EsP both start with EDUKASYON
    ('mapeh_tle', ('MAPEH', 'MUSIC', 'ARTS', 'PHYSICAL EDUCATION', 'HEALTH', 'TECHNOLOGY AND LIVELIHOOD', 'TLE', 'EPP',
                   'PANTAHANAN', 'PANGKABUHAYAN')),
    ('math_science', ('MATH', 'SCIENCE')),
    ('languages', ('ENGLISH', 'FILIPINO', 'ARALING', 'PAGPAPAKATAO', 'ESP', 'MOTHER TONGUE', 'MTB', 'GMRC',
                   'GOOD MANNERS', 'VALUES EDUCATION', 'MAKABANSA')),
)
MAPEH_PARTS = ('MUSIC', 'ARTS', 'PE', 'HEALTH')
# Senior High: INPUT DATA "SUBJECT TYPE" -> weights ('Helper(IMPORTANT!)'!CH2:CN6 of the SSHS template)
SSHS_TYPES = {
    'Core Subject (All Tracks)': (Decimal('0.25'), Decimal('0.5'), Decimal('0.25')),
    'Academic Elective (All Other Electives)': (Decimal('0.25'), Decimal('0.45'), Decimal('0.3')),
    'Academic Elective (Field Experience / Exposure, and Sports and Arts)': (Decimal('0.2'), Decimal('0.6'), Decimal('0.2')),
    'TechPro Elective (All Other Electives)': (Decimal('0.15'), Decimal('0.65'), Decimal('0.2')),
    'TechPro Elective (Work Immersion)': (Decimal('0.2'), Decimal('0.8'), Decimal('0')),
}
SSHS_LRN_COLUMNS = ('AK', 'AL')     # hidden INPUT DATA columns of ARMS-made SSHS records: male, female LRNs
SSHS_FIRST_ROW = 11                 # INPUT DATA learner rows 11-60 (males in L, females in O)
SSHS_MALE_ROWS = range(13, 63)      # the same learners on each quarter sheet
SSHS_FEMALE_ROWS = range(64, 114)
_ORDINAL = {'FIRST': 1, 'SECOND': 2, 'THIRD': 3, 'FOURTH': 4, '1ST': 1, '2ND': 2, '3RD': 3, '4TH': 4}
_SUFFIXES = {'JR', 'SR', 'II', 'III', 'IV', 'V'}
DESCRIPTORS = ((90, 'Outstanding'), (85, 'Very_Satisfactory'), (80, 'Satisfactory'), (75, 'Fairly_Satisfactory'),
               (0, 'Did_Not_Meet_Expectations'))

# ---- the three-term record of DepEd Order No. 015, s. 2026 (Grades 2 to 10)
# HELPER!B8:D48 of the template: lowest initial grade of each band -> term grade (the adjusted table of SY 2026-2027)
TERM_TRANSMUTATION = tuple((Decimal(str(low)), grade) for low, grade in (
    (99.5, 100), (98.32, 99), (97.14, 98), (95.96, 97), (94.78, 96), (93.6, 95), (92.42, 94), (91.24, 93), (90.06, 92),
    (88.88, 91), (87.7, 90), (86.52, 89), (85.34, 88), (84.16, 87), (82.98, 86), (81.8, 85), (80.62, 84), (79.44, 83),
    (78.26, 82), (77.08, 81), (75.9, 80), (74.72, 79), (73.54, 78), (72.36, 77), (71.18, 76), (70, 75), (65.34, 74),
    (60.67, 73), (56.01, 72), (51.34, 71), (46.67, 70), (42.01, 69), (37.34, 68), (32.68, 67), (28.01, 66), (23.35, 65),
    (18.68, 64), (14.01, 63), (9.35, 62), (4.68, 61), (0, 60)))
# Component weights for Grades 4 to 10 (written/oral works, product/performance tasks, examinations)
TERM_AREA_WEIGHTS = {
    'core': ('Languages, Mathematics, Science, Araling Panlipunan and GMRC/Values Education',
             (Decimal('0.2'), Decimal('0.5'), Decimal('0.3'))),
    'mapeh_tle': ('MAPEH and EPP/TLE', (Decimal('0.2'), Decimal('0.6'), Decimal('0.2'))),
}
TERM_FIRST_ROW = 11                 # INPUT DATA learner rows 11-60: males in K, females in N
TERM_LRN_COLUMNS = ('AK', 'AL')     # hidden INPUT DATA columns of ARMS-made three-term records: male, female LRNs
TERM_ICT_SHARE = Decimal('0.25')    # the TLE record kept per component: ICT counts 25%, the term's other area 75%
TERM_MAPEH_PARTS = ('MA', 'PEH')    # Music and Arts; Physical Education and Health


class EcrError(Exception):
    """The file cannot be used at all (not an e-class record, wrong period, unreadable)."""


# ---------------------------------------------------------------------------------------------
# arithmetic exactly as the template computes it
# ---------------------------------------------------------------------------------------------

def dec(value):
    """A cell value as a Decimal, or None when the cell is empty or not a number."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return Decimal(repr(value))
    text = str(value).strip()
    if not text:
        return None
    try:
        return Decimal(text)
    except Exception:
        return None


def xround(value, places=2):
    """Excel ROUND: halves go away from zero."""
    return None if value is None else value.quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)


def transmute(initial):
    """VLOOKUP(initial grade, TRANSMUTATION_TABLE, 4, TRUE)."""
    if initial is None:
        return None
    grade = None
    for low, q in TRANSMUTATION:
        if initial >= low:
            grade = q
        else:
            break
    return grade


def descriptor(quarterly):
    return next(name for low, name in DESCRIPTORS if quarterly >= low)


def learning_area(*names):
    """'languages', 'math_science' or 'mapeh_tle' from a subject or discipline name, else None."""
    text = ' '.join(str(n or '') for n in names).upper()
    for key, words in _AREA_WORDS:
        if any(re.search(r'\b' + re.escape(w), text) for w in words):
            return key
    return None


def compute(scores, weights, every_part=False):
    """One learner's grade from the raw scores, following the template cell by cell.

    scores: {'ww': [..], 'pt': [..], 'qa': [..]} item scores (Decimal or None) and 'hps' with the same keys.
    every_part: the Senior High rule. Its initial grade is WS(written work) + WS(performance tasks)
    + WS(quarterly assessment), so it needs a score in every component that carries weight (Work
    Immersion has none for the quarterly assessment). Grades 7 to 10 grade a learner once there is
    a quarterly assessment score, a blank component counting as 0.
    Returns totals, percentage and weighted scores, the initial and the quarterly grade, and the
    components still without a score ('missing').
    """
    out = {}
    for part, weight in zip(('ww', 'pt', 'qa'), weights):
        items = [s for s in scores[part] if s is not None]
        total = sum(items) if items else None                         # IF(COUNT(..)=0,"",SUM(..))
        hps = [h for h in scores['hps'][part] if h is not None]
        hps_total = sum(hps) if hps else None
        ps = xround(total / hps_total * 100) if total is not None and hps_total else None
        ws = xround(ps * weight) if ps is not None else None
        out[part] = {'total': total, 'hps': hps_total, 'ps': ps, 'ws': ws}
    if every_part:
        needed = [p for p, w in zip(('ww', 'pt', 'qa'), weights) if w]
        out['missing'] = [p for p in needed if out[p]['ws'] is None]
        initial = None if out['missing'] else xround(sum(out[p]['ws'] for p in needed))
    else:
        has_qa = any(s is not None for s in scores['qa'])
        out['missing'] = [] if has_qa else ['qa']
        # the template only grades a learner who has a quarterly assessment score; blank components count as 0
        initial = xround(sum(out[p]['ws'] or Decimal(0) for p in ('ww', 'pt', 'qa'))) if has_qa else None
    out['initial'] = initial
    out['quarterly'] = transmute(initial)
    return out


# ---------------------------------------------------------------------------------------------
# reading the workbook
# ---------------------------------------------------------------------------------------------

def _check_size(uploaded):
    size = getattr(uploaded, 'size', None)
    if size is not None and size > MAX_UPLOAD_BYTES:
        raise EcrError(f'The file is {size // (1024 * 1024)} MB. E-class records are about 2 MB; files up to 15 MB are accepted.')
    data = uploaded.read()
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED_BYTES:
                raise EcrError('The workbook is unusually large inside. Save it again from Excel and retry.')
    except zipfile.BadZipFile:
        raise EcrError('This is not an Excel workbook (.xlsx). Save the e-class record as .xlsx and upload it again.')
    return data


def _grid(ws):
    """All values of a sheet as rows of a list, 1-based through padding."""
    rows = [[None]]
    for row in ws.iter_rows(values_only=True):
        rows.append([None] + list(row))
    return rows


def _cell(grid, row, col):
    try:
        return grid[row][col]
    except IndexError:
        return None


def _text(value):
    return '' if value is None or value == 0 else str(value).strip()


def _col(letter):
    n = 0
    for ch in letter.upper():
        n = n * 26 + ord(ch) - 64
    return n


def _period_number(sheet_name, label):
    """1 for 'FIRST QUARTER' / 'MATH_Q1' / 'SCIENCE _T1', and so on."""
    words = re.findall(r'[A-Z0-9]+', (label or '').upper())
    for w in words:
        if w in _ORDINAL:
            return _ORDINAL[w]
    m = re.search(r'(?:^|[_\s])[QTS]\s*([1-4])\s*$', sheet_name.upper())
    return int(m.group(1)) if m else None


def _layout(grid, header=8, sub=9):
    """Find the component columns from the headers instead of trusting fixed positions."""
    groups = {}
    for col in range(1, len(grid[header]) if len(grid) > header else 0):
        text = str(_cell(grid, header, col) or '').strip().upper()
        if 'WRITTEN WORK' in text:
            groups['ww'] = col
        elif 'PERFORMANCE TASK' in text:
            groups['pt'] = col
        elif 'QUARTERLY ASSESSMENT' in text:
            groups['qa'] = col
        elif text.startswith('INITIAL'):
            groups['initial'] = col
        elif text.startswith('QUARTERLY'):
            groups['quarterly'] = col
    if not {'ww', 'pt', 'qa', 'initial', 'quarterly'} <= set(groups):
        return _layout_from_numbers(grid, sub)
    layout = {'initial': groups['initial'], 'quarterly': groups['quarterly']}
    bounds = sorted((groups[p], p) for p in ('ww', 'pt', 'qa')) + [(groups['initial'], 'end')]
    for (start, part), (nxt, _) in zip(bounds, bounds[1:]):
        items, ps_col, ws_col = [], None, None
        for col in range(start, nxt):
            label = str(_cell(grid, sub, col) or '').strip().upper()
            if label == 'PS':
                ps_col = col
            elif label == 'WS':
                ws_col = col
            elif label not in ('TOTAL', '') and ps_col is None:
                items.append(col)
        if not items or ws_col is None:
            return None
        layout[part] = {'items': items, 'weight_col': ws_col}
    return layout


def _layout_from_numbers(grid, sub):
    """The same columns from the numbered sub-header alone (1..10, Total, PS, WS for each component).

    The Senior High record writes its component headers with formulas that stay blank until a
    subject type is chosen; its numbered sub-header never changes.
    """
    if len(grid) <= sub:
        return None
    labels = [(col, str(v).strip().upper() if v is not None else '') for col, v in enumerate(grid[sub]) if col]
    ws_cols = [col for col, text in labels if text == 'WS']
    if len(ws_cols) < 3:
        return None
    layout, start = {}, 1
    for part, ws_col in zip(('ww', 'pt', 'qa'), ws_cols[:3]):
        items = [col for col, text in labels if start <= col < ws_col and re.fullmatch(r'\d+(\.0)?', text)]
        if not items:
            return None
        layout[part] = {'items': items, 'weight_col': ws_col}
        start = ws_col + 1
    layout['initial'], layout['quarterly'] = ws_cols[2] + 1, ws_cols[2] + 2
    return layout


def _weights(grid, layout, row=10):
    return tuple(dec(_cell(grid, row, layout[p]['weight_col'])) for p in ('ww', 'pt', 'qa'))


def subject_type_weights(text):
    """Weights of an SSHS subject type as written in INPUT DATA, or None."""
    key = _clean(text)
    return next((w for name, w in SSHS_TYPES.items() if _clean(name) == key), None) if key else None


def _sshs_meta(inp):
    """INPUT DATA of the SSHS record: each label in column C, its value in F."""
    labels = {'REGION': 'region', 'DIVISION': 'division', 'SCHOOL ID': 'school_id', 'SCHOOL NAME': 'school',
              'SCHOOL YEAR': 'school_year', 'TEACHER': 'teacher', 'GRADE LEVEL': 'grade', 'SECTION NAME': 'section_name',
              'SUBJECT TYPE': 'subject_type', 'SUBJECT': 'subject_listed'}
    meta = {key: '' for key in labels.values()}
    meta['subject_other'] = ''
    for r in range(1, min(len(inp), 60)):
        label = _text(_cell(inp, r, 3)).upper()
        value = _cell(inp, r, _col('F'))
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        if label in labels:
            meta[labels[label]] = _text(value)
        elif label.startswith('OTHER ELECTIVE'):
            meta['subject_other'] = _text(value)
    if meta['region'].upper().startswith('SELECT'):
        meta['region'] = ''
    listed = meta['subject_listed']
    if listed.upper().startswith(('OTHER ELECTIVE AND SPECIAL', 'INPUT GRADE')):
        listed = ''
    meta['subject'] = meta['subject_other'] or listed
    grade = meta['grade']
    meta['section'] = f'Grade {grade} - {meta["section_name"]}' if grade and meta['section_name'] else meta['section_name']
    return meta


def _is_sshs(inp):
    return any(_text(_cell(inp, r, 3)).upper() == 'SUBJECT TYPE' for r in range(1, min(len(inp), 60)))


def read_workbook(uploaded):
    """Every grading-period sheet of an e-class record with its learners' scores.

    Returns {'kind': 'JHS' | 'SSHS' | 'TERM', 'meta': {...}, 'periods': {number: [sheet, ...]}} where each
    sheet carries its layout, weights, highest possible scores and rows (sex, name, LRN if ARMS made
    the file, item scores and the grade the file itself shows).
    """
    data = _check_size(uploaded)
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise EcrError('The workbook could not be opened. Save it again from Excel as .xlsx and retry.')
    try:
        input_name = next((n for n in wb.sheetnames if 'INPUT' in n.upper()), None)
        if not input_name:
            raise EcrError('This file has no INPUT DATA sheet, so it is not a DepEd e-class record.')
        inp = _grid(wb[input_name])
        if _is_term(inp):
            return _read_terms(wb, inp, input_name)
        sshs = _is_sshs(inp)
        if sshs:
            meta = _sshs_meta(inp)
            names, lrns = {}, {}
            for rows, name_col, lrn_col in ((SSHS_MALE_ROWS, 'L', SSHS_LRN_COLUMNS[0]), (SSHS_FEMALE_ROWS, 'O', SSHS_LRN_COLUMNS[1])):
                for i, r in enumerate(rows):
                    names[r] = _text(_cell(inp, SSHS_FIRST_ROW + i, _col(name_col)))
                    lrns[r] = re.sub(r'\D', '', str(_cell(inp, SSHS_FIRST_ROW + i, _col(lrn_col)) or ''))
            header, hps_row, marker_cols = 9, 11, (1, 2)
        else:
            meta = {
                'region': _text(_cell(inp, 4, _col('G'))), 'division': _text(_cell(inp, 4, _col('O'))),
                'school': _text(_cell(inp, 5, _col('G'))), 'school_id': _text(_cell(inp, 5, _col('X'))),
                'school_year': _text(_cell(inp, 5, _col('AG'))), 'section': _text(_cell(inp, 7, _col('K'))),
                'teacher': _text(_cell(inp, 7, _col('S'))), 'subject': _text(_cell(inp, 7, _col('AG'))),
            }
            names = {r: _text(_cell(inp, r, 2)) for r in range(12, len(inp))}
            lrns = {r: re.sub(r'\D', '', str(_cell(inp, r, _col(LRN_COLUMN)) or '')) for r in range(12, len(inp))}
            header, hps_row, marker_cols = 8, 10, (2,)

        periods = {}
        for name in wb.sheetnames:
            if name == input_name or 'SUMMARY' in name.upper() or 'DO NOT DELETE' in name.upper():
                continue
            ws = wb[name]
            if getattr(ws, 'sheet_state', 'visible') == 'veryHidden':
                continue
            grid = _grid(ws)
            if not any('HIGHEST POSSIBLE' in str(_cell(grid, hps_row, c) or '').upper() for c in (1, 2)):
                continue
            layout = _layout(grid, header, header + 1)
            if not layout:
                raise EcrError(f'Sheet "{name}" does not have the e-class record columns '
                               '(written works, performance tasks, quarterly assessment). Use the official template.')
            label = _text(_cell(grid, 7, 1))
            number = _period_number(name, label)
            if not number:
                raise EcrError(f'Sheet "{name}" does not say which grading period it is (cell A7 is "{label}").')
            hps = {p: [dec(_cell(grid, hps_row, c)) for c in layout[p]['items']] for p in ('ww', 'pt', 'qa')}
            rows, sex = [], None
            for r in range(hps_row + 1, len(grid)):
                marker = next((m for m in (_text(_cell(grid, r, c)).upper() for c in marker_cols) if m in ('MALE', 'FEMALE')), '')
                if marker:
                    sex = marker[0]
                    continue
                if sex is None or not isinstance(_cell(grid, r, 1), (int, float)):
                    continue                       # learner rows are the numbered rows under MALE / FEMALE
                raw = {p: [_cell(grid, r, c) for c in layout[p]['items']] for p in ('ww', 'pt', 'qa')}
                name_text = _text(_cell(grid, r, 2)) or names.get(r, '')
                filled = [v for vals in raw.values() for v in vals if v not in (None, '')]
                if not name_text and not filled:
                    continue
                rows.append({
                    'row': r, 'sex': sex, 'name': name_text, 'lrn': lrns.get(r, ''), 'raw': raw,
                    'file_initial': dec(_cell(grid, r, layout['initial'])),
                    'file_quarterly': dec(_cell(grid, r, layout['quarterly'])),
                })
            if sshs:
                component = ''
                weights = _weights(grid, layout, hps_row)
                if None in weights:                  # formulas never calculated: the subject type decides
                    weights = subject_type_weights(meta['subject_type']) or weights
            else:
                component = _text(_cell(grid, 7, _col('AG'))) or re.sub(r'[_\s]*[QTS][1-4]$', '', name.strip(), flags=re.I)
                weights = _weights(grid, layout, hps_row)
            periods.setdefault(number, []).append({
                'sheet': name, 'label': label, 'component': component.strip(), 'layout': layout,
                'weights': weights, 'hps': hps, 'hps_row': hps_row, 'rows': rows,
            })
    finally:
        wb.close()
    if not periods:
        raise EcrError('No grading-period sheet was found. Upload the e-class record (.xlsx) made from the DepEd template.')
    return {'kind': 'SSHS' if sshs else 'JHS', 'meta': meta, 'periods': periods}


# ---------------------------------------------------------------------------------------------
# learners: names in the file against the class list
# ---------------------------------------------------------------------------------------------

def _clean(text):
    text = unicodedata.normalize('NFKD', str(text or '')).encode('ascii', 'ignore').decode()
    text = text.replace("'", '').replace('-', ' ').replace('.', ' ')
    text = re.sub(r'[^A-Za-z0-9, ]+', ' ', text)
    return re.sub(r'\s+', ' ', text).strip().upper()


def _words(text):
    return [w for w in _clean(text).replace(',', ' ').split() if w not in _SUFFIXES]


def _compact(words):
    return ''.join(words)


def name_score(file_name, person):
    """How surely a name in the file is this learner: 100 exact, 90 or more is certain, below needs a person."""
    first, middle, last = _words(person['first']), _words(person['middle']), _words(person['last'])
    cleaned = _clean(file_name)
    if not cleaned or not first or not last:
        return 0
    if ',' in cleaned:                                   # LAST, FIRST MIDDLE (the DepEd convention)
        f_last, rest = cleaned.split(',', 1)
        f_last, rest = _words(f_last), _words(rest)
        if _compact(f_last) == _compact(last) and rest[:len(first)] == first:
            tail = rest[len(first):]
            if not tail:
                return 95
            if _compact(tail) == _compact(middle) or (middle and len(tail) == 1 and tail[0] == middle[0][0]) or \
                    (middle and [t[0] for t in tail] == [m[0] for m in middle]):
                return 100
            return 70                                    # same first and last name, different middle name
    else:                                                # First Middle Last, in any order
        tokens = set(_words(cleaned))
        needed = set(first) | set(last)
        allowed = needed | set(middle) | {m[0] for m in middle}
        if needed <= tokens <= allowed:
            return 95 if tokens & (set(middle) | {m[0] for m in middle}) or not middle else 92
        if _compact(_words(cleaned)) in (_compact(first + last), _compact(first + middle + last)):
            return 92
    a = ' '.join(last + first)
    b = ' '.join(_words(cleaned.replace(',', ' ')))
    b_alt = ' '.join(_words(cleaned.split(',', 1)[1] + ' ' + cleaned.split(',', 1)[0])) if ',' in cleaned else b
    c = ' '.join(first + last)
    ratio = max(SequenceMatcher(None, a, b).ratio(), SequenceMatcher(None, c, b).ratio(),
                SequenceMatcher(None, c, b_alt).ratio(), SequenceMatcher(None, a, b_alt).ratio())
    return min(89, int(ratio * 100))


def match_rows(rows, roster):
    """Pair every file row with one learner of the class, or say why it cannot be done safely.

    roster: [{'id' (enrollment), 'lrn', 'first', 'middle', 'last', 'sex', 'active'}]
    Each row gets 'status': 'matched' (certain), 'check' (a person must confirm: sex differs, the LRN
    and the name disagree, a learner appears twice, the learner left the class), 'suggested' (close
    spelling only) or 'unmatched'; plus 'enrollment' and 'candidates' (best first).
    """
    by_lrn = {p['lrn']: p for p in roster if p.get('lrn')}
    for row in rows:
        scored = sorted(((name_score(row['name'], p), p) for p in roster), key=lambda sp: -sp[0])
        row['candidates'] = [{'id': p['id'], 'score': s} for s, p in scored[:3] if s >= 60]
        row['reasons'] = []
        best_score, best = scored[0] if scored else (0, None)
        runner = scored[1][0] if len(scored) > 1 else 0
        lrn_person = by_lrn.get(row.get('lrn')) if row.get('lrn') else None

        if lrn_person:
            row['enrollment'] = lrn_person['id']
            if name_score(row['name'], lrn_person) >= 90:
                row['status'] = 'matched'
            else:
                row['status'] = 'check'
                row['reasons'].append('The LRN in the file belongs to a learner with a different name.')
        elif best and best_score >= 90 and runner < 90:
            row['enrollment'], row['status'] = best['id'], 'matched'
            if row.get('lrn'):
                row['status'] = 'check'
                row['reasons'].append('The LRN in the file is not in this class.')
        elif best and best_score >= 90:
            row['enrollment'], row['status'] = None, 'check'
            row['reasons'].append('More than one learner in the class has this name.')
        elif best and best_score >= 75:
            row['enrollment'], row['status'] = best['id'], 'suggested'
            row['reasons'].append('Only a similar spelling was found.')
        else:
            row['enrollment'], row['status'] = None, 'unmatched'
            row['reasons'].append('No learner of this class has this name.')

        person = next((p for p in roster if p['id'] == row['enrollment']), None)
        if person and row['status'] == 'matched':
            if person['sex'] and row['sex'] and person['sex'] != row['sex']:
                row['status'] = 'check'
                row['reasons'].append(f'Listed under {"MALE" if row["sex"] == "M" else "FEMALE"} in the file, '
                                      f'but the learner is recorded as {"male" if person["sex"] == "M" else "female"}.')
            if not person.get('active', True):
                row['status'] = 'check'
                row['reasons'].append('This learner is no longer enrolled in the class.')

    # one learner, one row: a second claim needs a person to decide
    claimed = {}
    for row in rows:
        if row.get('enrollment') and row['status'] in ('matched', 'check', 'suggested'):
            claimed.setdefault(row['enrollment'], []).append(row)
    for twins in claimed.values():
        if len(twins) > 1:
            for row in twins:
                row['status'] = 'check'
                row['reasons'].append('Two rows of the file point to the same learner.')
    return rows


# ---------------------------------------------------------------------------------------------
# one period of one class, ready to show and to save
# ---------------------------------------------------------------------------------------------

def _scores(raw, hps, prefix='', hps_row=10):
    """Item scores as Decimals, with the reasons any score cannot be used."""
    errors, scores = [], {'hps': hps}
    labels = {'ww': 'Written work', 'pt': 'Performance task', 'qa': 'Quarterly assessment'}
    for part in ('ww', 'pt', 'qa'):
        values = []
        for i, (value, top) in enumerate(zip(raw[part], hps[part]), start=1):
            if value in (None, ''):
                values.append(None)
                continue
            number = dec(value)
            name = prefix + (f'{labels[part]} {i}' if part != 'qa' else labels[part])
            if number is None:
                errors.append(f'{name}: "{value}" is not a number.')
            elif number < 0:
                errors.append(f'{name}: {number} is below zero.')
            elif top is None or top <= 0:
                errors.append(f'{name}: has a score but no highest possible score in row {hps_row}.')
            elif number > top:
                errors.append(f'{name}: {number} is more than the highest possible score of {top}.')
            values.append(number if number is not None and number >= 0 and top and top > 0 and number <= top else None)
        scores[part] = values
    return scores, errors


def analyse(parsed, period_number, *, subject_names, senior_high, roster, mapeh=False, subject_type=None):
    """Grades of one grading period from a parsed workbook, checked and matched to the class.

    subject_names: the class subject and discipline names (DepEd weights check).
    subject_type: for a Senior High class, the SSHS subject type ARMS expects (a mismatch is a warning).
    Returns {'sheets', 'checks', 'rows', 'other_periods'}. checks are (level, text) with level
    'error' (the file cannot be imported) or 'warn'. Every row has:
      quarterly / initial  the grade the scores give (None when the template gives no grade)
      errors               reasons the row cannot be imported (a score is unusable, scores without a name)
      warnings             things the teacher must confirm (the file shows another grade)
      notes                plain facts (no quarterly assessment yet, so no grade)
      status, enrollment, candidates, reasons   from match_rows
    """
    if parsed.get('kind') == 'TERM':
        return _analyse_terms(parsed, period_number, subject_names=subject_names, roster=roster, mapeh=mapeh)
    sheets = parsed['periods'].get(period_number)
    if not sheets:
        found = ', '.join(str(n) for n in sorted(parsed['periods']))
        raise EcrError(f'The file has no sheet for grading period {period_number} (it has period {found}).')
    if mapeh:
        parts = {}
        for s in sheets:
            key = mapeh_part(s['component'] or s['sheet'])
            if key:
                parts[key] = s
        missing = [p for p in MAPEH_PARTS if p not in parts]
        if missing:
            raise EcrError('A MAPEH grade needs Music, Arts, PE and Health. The file has no '
                           f'{", ".join(missing)} sheet for this period.')
        sheets = [parts[p] for p in MAPEH_PARTS]
    elif len(sheets) > 1:
        raise EcrError('The file has more than one subject for this period. Upload one subject per file '
                       '(the four MAPEH sheets are read together when the class is MAPEH).')

    checks = []
    sshs = parsed.get('kind') == 'SSHS'
    area = None if senior_high or sshs else learning_area(*subject_names)
    for s in sheets:
        w = s['weights']
        row = s.get('hps_row', 10)
        if sshs and None in w:
            checks.append(('error', f'{s["sheet"]}: no subject type is chosen in INPUT DATA (SUBJECT TYPE), so the record '
                                    'has no weights. Choose it, save the file in Excel and upload it again.'))
        elif None in w or abs(sum(w) - 1) > Decimal('0.0001'):
            checks.append(('error', f'{s["sheet"]}: the component weights in row {row} do not add up to 100%.'))
        elif sshs:
            if tuple(w) not in SSHS_TYPES.values():
                checks.append(('error', f'{s["sheet"]} weighs written work, performance tasks and quarterly assessment '
                                        f'{_pct(w)}, which is not one of the SSHS subject types. The formulas may have been changed.'))
            elif subject_type and SSHS_TYPES.get(subject_type) != tuple(w):
                checks.append(('warn', f'The file grades this subject as {parsed["meta"].get("subject_type") or _pct(w)} '
                                       f'({_pct(w)}); ARMS lists it as {subject_type} ({_pct(SSHS_TYPES[subject_type])}). '
                                       'The file\'s weights are used.'))
        elif area:
            label, expected = AREA_WEIGHTS[area]
            if w != expected:
                checks.append(('error', f'{s["sheet"]} weighs written work, performance tasks and quarterly assessment '
                                        f'{_pct(w)}, but {label} use {_pct(expected)} (DepEd Order No. 8, s. 2015). '
                                        'This looks like the template of another subject.'))
        for (p, name), weight in zip((('ww', 'written work'), ('pt', 'performance tasks'), ('qa', 'quarterly assessment')),
                                     w if None not in w else (1, 1, 1)):
            if weight and not any(h for h in s['hps'][p] if h):
                checks.append(('warn', f'{s["sheet"]}: there is no highest possible score for {name} in row {row}.'))

    # every component sheet lists the same learners in the same rows (they all read INPUT DATA)
    rows = []
    for base in sheets[0]['rows']:
        entry = {'row': base['row'], 'sex': base['sex'], 'name': base['name'], 'lrn': base['lrn'],
                 'errors': [], 'warnings': [], 'notes': [], 'parts': []}
        for s in sheets:
            same = next((x for x in s['rows'] if x['row'] == base['row']), None)
            prefix = f'{s["component"]}: ' if mapeh else ''
            if same is None:
                entry['parts'].append({'component': s['component'], 'result': compute({'ww': [], 'pt': [], 'qa': [], 'hps': s['hps']}, s['weights']),
                                       'has_scores': False})
                continue
            scores, errors = _scores(same['raw'], s['hps'], prefix, s.get('hps_row', 10))
            weights = tuple(w if w is not None else Decimal(0) for w in s['weights'])
            result = compute(scores, weights, every_part=sshs)
            entry['errors'] += errors
            has_scores = any(v not in (None, '') for vals in same['raw'].values() for v in vals)
            if has_scores and result['quarterly'] is None and not errors:
                missing = ', '.join(_PART_NAMES[p] for p in result['missing']) or 'score'
                entry['notes'].append(f'{prefix}no {missing} score yet, so the template gives no grade.')
            if same['file_quarterly'] is not None and result['quarterly'] is not None and (
                    same['file_quarterly'] != result['quarterly'] or
                    (same['file_initial'] is not None and abs(same['file_initial'] - result['initial']) > Decimal('0.01'))):
                entry['warnings'].append(f'{prefix}the file shows {same["file_quarterly"]}, but its scores give '
                                         f'{result["quarterly"]} (initial grade {result["initial"]}). The formulas may have been changed.')
            entry['parts'].append({'component': s['component'], 'result': result, 'has_scores': has_scores})
        rows.append(entry)

    for entry in rows:
        results = [p['result'] for p in entry['parts']]
        entry['has_scores'] = any(p['has_scores'] for p in entry['parts'])
        if mapeh:
            graded = [r for r in results if r['quarterly'] is not None]
            if len(graded) == 4:
                entry['initial'] = xround(sum(r['initial'] for r in graded) / 4)
                entry['quarterly'] = int(xround(Decimal(sum(r['quarterly'] for r in graded)) / 4, 0))   # ROUND(AVERAGE(...),0)
            else:
                entry['initial'] = entry['quarterly'] = None
                if graded:
                    entry['notes'].append(f'Only {len(graded)} of the 4 MAPEH components have a grade, so there is no MAPEH grade yet.')
        else:
            entry['initial'], entry['quarterly'] = results[0]['initial'], results[0]['quarterly']
        if entry['errors']:
            entry['initial'] = entry['quarterly'] = None          # a score cannot be trusted, so neither can the grade
        if entry['has_scores'] and not entry['name']:
            entry['errors'].append('There are scores in this row but no learner name.')

    named = [e for e in rows if e['name']]
    match_rows(named, roster)
    for e in rows:
        if not e['name']:
            e.update(status='unmatched', enrollment=None, candidates=[], reasons=[])

    other = sorted(n for n, ss in parsed['periods'].items() if n != period_number and
                   any(any(v not in (None, '') for vals in r['raw'].values() for v in vals) for s in ss for r in s['rows']))
    return {'sheets': [s['sheet'] for s in sheets], 'checks': checks, 'rows': rows, 'other_periods': other,
            'weights': sheets[0]['weights']}


_PART_NAMES = {'ww': 'written work', 'pt': 'performance task', 'qa': 'quarterly assessment'}


def mapeh_part(name):
    """'MUSIC', 'ARTS', 'PE' or 'HEALTH' from a MAPEH sheet or component name ('Physical Education' is PE)."""
    text = _clean(name)
    if text.startswith('MUSIC'):
        return 'MUSIC'
    if text.startswith('ARTS') or text == 'ART':
        return 'ARTS'
    if text.startswith('PE') or 'PHYSICAL' in text:
        return 'PE'
    if text.startswith('HEALTH'):
        return 'HEALTH'
    return None


def _pct(weights):
    return '/'.join(str(int(w * 100)) for w in weights)


def same_section(file_section, section_name, grade_number):
    """Does the GRADE & SECTION cell name this class? Empty cells are not held against the file."""
    text = _clean(file_section)
    if not text:
        return True
    words = set(text.replace(',', ' ').split())
    wanted = _clean(section_name).replace(',', ' ').split()
    if not set(wanted) <= words:
        return False
    numbers = {int(n) for n in re.findall(r'\d+', text)}
    return not numbers or grade_number in numbers


_SUBJECT_KEYS = (       # Grades 7 to 10 learning areas and the ways their names are written
    ('MATH', r'\bMATH'),
    ('SCIENCE', r'\bSCIENCE\b'),
    ('ENGLISH', r'\bENGLISH\b'),
    ('FILIPINO', r'\bFILIPINO\b'),
    ('AP', r'\bARALING PANLIPUNAN\b|^AP\b'),
    ('ESP', r'\bEDUKASYON SA PAGPAPAKATAO\b|^ESP\b|\bVALUES EDUCATION\b|\bGMRC\b|\bGOOD MANNERS\b'),
    ('MAPEH', r'\bMAPEH\b|\bMUSIC\b|\bARTS?\b|\bPHYSICAL EDUCATION\b|\bHEALTH\b|^PE\b'),
    ('TLE', r'\bTECHNOLOGY AND LIVELIHOOD\b|\bTLE\b|\bEPP\b|\bEDUKASYONG PANTAHANAN\b'),
    ('MTB', r'\bMOTHER TONGUE\b|\bMTB\b'),
)


def subject_key(name):
    """MATH, SCIENCE, ENGLISH, ... for a Grades 7 to 10 learning area name, else None."""
    text = _clean(name)
    return next((key for key, pattern in _SUBJECT_KEYS if re.search(pattern, text)), None)


def same_subject(file_subject, *names, junior=True):
    """Does the SUBJECT cell name this class's subject (or its discipline)? An empty cell is not held against the file.

    Grades 7 to 10 compare learning areas (TLE is Technology and Livelihood Education); Senior High
    subjects compare their words, since several of them share a learning area.
    """
    text = _clean(file_subject)
    if not text:
        return True
    key = subject_key(text) if junior else None
    for name in names:
        n = _clean(name)
        if not n:
            continue
        if key and key == subject_key(n):
            return True
        if not key and (n in text or text in n):
            return True
    return False


# ---------------------------------------------------------------------------------------------
# the three-term record (DepEd Order No. 015, s. 2026)
# ---------------------------------------------------------------------------------------------
# A term sheet groups its score columns in blocks, each ending in Total, PS and WS on the numbered sub-header row: one
# block per component, or (GMRC and Values Education) one per domain, cognitive, affective and behavioral, each with
# its own weight. The examinations block holds Summative Test 1, Summative Test 2 and the Term Examination, which
# count 30, 30 and 40 of its 100. No cell is rounded until the term grade, and ARMS computes the same way.

_EXAM_NAMES = ('Summative test 1', 'Summative test 2', 'Term examination')
_TERM_PART_NAMES = {'ww': 'Written/oral work', 'pt': 'Performance task'}


def term_transmute(initial):
    """The term grade of an initial grade, as the record's own formula finds it:
    INDEX(HELPER!D8:D48, IF(IG >= 99.5, 1, MATCH(IG, HELPER!B8:B48, -1) + 1)).

    MATCH(..., -1) finds the smallest band minimum at or above the initial grade and the band is the row after it, so
    an initial grade exactly on a band's minimum lands a band lower (70.00 gives 74, not 75) and 0 gives no grade. The
    template does this, so ARMS does too: the grade a file shows and the one ARMS computes always agree.
    """
    if initial is None:
        return None
    if initial >= TERM_TRANSMUTATION[0][0]:
        return TERM_TRANSMUTATION[0][1]
    at = None
    for i, (low, _) in enumerate(TERM_TRANSMUTATION):
        if low >= initial:
            at = i
        else:
            break
    if at is None or at + 1 >= len(TERM_TRANSMUTATION):
        return None
    return TERM_TRANSMUTATION[at + 1][1]


def compute_term(blocks, values):
    """One learner's term grade from the raw scores, cell by cell as the three-term record computes it.

    blocks: the sheet's score blocks; values: one list of item scores (Decimal or None) per block. A block's percentage
    is its total over its highest possible score (the examinations: each test's share of its own highest score, 30, 30
    and 40), and its weighted score that times the block weight. The initial grade adds the weighted scores (a blank
    block counts 0) once the learner has any score. Returns, per component ('ww', 'pt', 'qa'), the totals, percentage
    and weighted score (a component of several blocks: their weighted average), the initial grade and the term grade.
    """
    out = {p: {'total': None, 'hps': None, 'ps': None, 'ws': None, 'weight': Decimal(0)} for p in ('ww', 'pt', 'qa')}
    initial, scored = Decimal(0), False
    for block, items in zip(blocks, values):
        present = [v for v in items if v is not None]
        scored = scored or bool(present)
        part = out[block['part']]
        part['weight'] += block['weight'] or Decimal(0)
        hps_total = sum(h for h in block['hps'] if h) or None
        total = sum(present) if present else None
        if block['tests']:
            shares = [v / h * t for v, h, t in zip(items, block['hps'], block['tests']) if v is not None and h and t is not None]
            ps = sum(shares) if shares else None
        else:
            ps = total / hps_total * 100 if total is not None and hps_total else None
        ws = ps * block['weight'] if ps is not None and block['weight'] is not None else None
        if total is not None:
            part['total'] = (part['total'] or Decimal(0)) + total
        if hps_total:
            part['hps'] = (part['hps'] or Decimal(0)) + hps_total
        if ws is not None:
            part['ws'] = (part['ws'] or Decimal(0)) + ws
            initial += ws
    for part in out.values():
        part['ps'] = xround(part['ws'] / part['weight']) if part['ws'] is not None and part['weight'] else None
        part['ws'] = xround(part['ws'])
    raw = initial if scored else None
    out.update(initial=xround(raw), quarterly=term_transmute(raw), missing=[])
    return out


def _is_term(inp):
    return any(_text(_cell(inp, r, 3)).upper() == 'SUBJECT TEACHER' for r in range(1, min(len(inp), 40)))


def _domain(grid, header, sub, col):
    """'cognitive', 'affective' or 'behavioral' when the block sits under a domain label (GMRC and Values Education)."""
    for r in range(header + 1, sub):
        text = _text(_cell(grid, r, col)).upper()
        if text.endswith('DOMAIN'):
            return text.replace('DOMAIN', '').strip().lower()
    return ''


def _term_layout(grid):
    """The score blocks of a term sheet, found from its headers and its numbered sub-header row."""
    header = next((r for r in range(1, min(len(grid), 30))
                   if any('WRITTEN' in str(v or '').upper() and 'ORAL' in str(v or '').upper() for v in grid[r])), None)
    if header is None:
        return None
    cols = {}
    for c, v in enumerate(grid[header]):
        t = str(v or '').strip().upper()
        if 'WRITTEN' in t and 'ww' not in cols:
            cols['ww'] = c
        elif 'PERFORMANCE' in t and 'pt' not in cols:
            cols['pt'] = c
        elif 'EXAMINATION' in t and 'qa' not in cols:
            cols['qa'] = c
        elif t.startswith('INITIAL'):
            cols['initial'] = c
        elif t.startswith('TERM GRADE'):
            cols['term'] = c
    if not {'ww', 'pt', 'qa', 'initial', 'term'} <= set(cols):
        return None
    hps_row = next((r for r in range(header + 1, min(len(grid), header + 8))
                    if 'HIGHEST POSSIBLE' in _text(_cell(grid, r, 2)).upper()), None)
    if hps_row is None:
        return None
    sub = hps_row - 1
    blocks, items, tests = [], [], []
    for c in range(cols['ww'], cols['initial']):
        label = _text(_cell(grid, sub, c)).upper()
        if re.fullmatch(r'\d+(\.0)?', label) or label in ('ST1', 'ST2', 'TE'):
            items.append(c)
        elif label.startswith('WS '):
            tests.append(dec(_cell(grid, hps_row, c)))
        elif label == 'WS':
            part = 'ww' if c < cols['pt'] else ('pt' if c < cols['qa'] else 'qa')
            if not items or (part == 'qa' and len(tests) != len(items)):
                return None
            blocks.append({'part': part, 'items': items, 'tests': tests if part == 'qa' else None,
                           'weight': dec(_cell(grid, hps_row, c)), 'hps': [dec(_cell(grid, hps_row, i)) for i in items],
                           'domain': _domain(grid, header, sub, items[0])})
            items, tests = [], []
    if {b['part'] for b in blocks} != {'ww', 'pt', 'qa'}:
        return None
    return {'blocks': blocks, 'initial': cols['initial'], 'term': cols['term'], 'hps_row': hps_row}


def _read_terms(wb, inp, input_name):
    """Every term sheet of a three-term record with its learners' scores (see read_workbook)."""
    labels = {'REGION': 'region', 'DIVISION': 'division', 'SCHOOL ID': 'school_id', 'SCHOOL NAME': 'school',
              'SCHOOL YEAR': 'school_year', 'SCHOOL HEAD': 'school_head', 'SUBJECT TEACHER': 'teacher', 'SUBJECT': 'subject',
              'GRADE LEVEL': 'grade', 'SECTION': 'section_name'}
    meta = {key: '' for key in labels.values()}
    for r in range(1, min(len(inp), 40)):
        key = labels.get(_text(_cell(inp, r, 3)).upper())
        if key:
            value = _cell(inp, r, _col('E'))
            meta[key] = _text(int(value) if isinstance(value, float) and value.is_integer() else value)
    grade = meta['grade']
    meta['section'] = f'{grade} - {meta["section_name"]}' if grade and meta['section_name'] else meta['section_name']
    names = {'M': [], 'F': []}
    lrns = {'M': [], 'F': []}
    for i in range(50):
        r = TERM_FIRST_ROW + i
        for sex, name_col, lrn_col in (('M', 'K', TERM_LRN_COLUMNS[0]), ('F', 'N', TERM_LRN_COLUMNS[1])):
            names[sex].append(_text(_cell(inp, r, _col(name_col))))
            lrns[sex].append(re.sub(r'\D', '', str(_cell(inp, r, _col(lrn_col)) or '')))

    periods = {}
    for name in wb.sheetnames:
        m = re.match(r'\s*TERM\s*([1-3])\b(.*)$', name, re.I)
        if name == input_name or not m or 'SUMMARY' in name.upper():
            continue
        ws = wb[name]
        if getattr(ws, 'sheet_state', 'visible') == 'veryHidden':
            continue
        grid = _grid(ws)
        layout = _term_layout(grid)
        if not layout:
            raise EcrError(f'Sheet "{name}" does not have the three-term e-class record columns (written/oral works, '
                           'product/performance tasks, examinations). Use the official template.')
        label = _text(_cell(grid, 10, 2))
        number = _period_number(name, label) or int(m.group(1))
        rows, sex, index = [], None, 0
        for r in range(layout['hps_row'] + 1, len(grid)):
            marker = _text(_cell(grid, r, 2)).upper()
            if marker in ('MALE', 'FEMALE'):
                sex, index = marker[0], 0
                continue
            if sex is None or not isinstance(_cell(grid, r, 2), (int, float)):
                continue                           # learner rows are the numbered rows under MALE / FEMALE
            i, index = index, index + 1
            raw = [[_cell(grid, r, c) for c in b['items']] for b in layout['blocks']]
            name_text = _text(_cell(grid, r, 3)) or (names[sex][i] if i < 50 else '')
            if not name_text and not any(v not in (None, '') for vals in raw for v in vals):
                continue
            rows.append({'row': r, 'sex': sex, 'name': name_text, 'lrn': lrns[sex][i] if i < 50 else '', 'raw': raw,
                         'file_initial': dec(_cell(grid, r, layout['initial'])),
                         'file_quarterly': dec(_cell(grid, r, layout['term']))})
        weights = tuple(sum((b['weight'] or Decimal(0)) for b in layout['blocks'] if b['part'] == p) for p in ('ww', 'pt', 'qa'))
        periods.setdefault(number, []).append({'sheet': name, 'label': label, 'component': m.group(2).strip(), 'layout': layout,
                                               'weights': weights, 'hps_row': layout['hps_row'], 'rows': rows})
    if not periods:
        raise EcrError('No term sheet was found. Upload the three-term e-class record (.xlsx) made from the DepEd template.')
    return {'kind': 'TERM', 'meta': meta, 'periods': periods}


def term_mapeh_part(name):
    """'MA' (Music and Arts) or 'PEH' (Physical Education and Health) from a three-term MAPEH sheet's name."""
    text = _clean(name)
    if text.startswith('M AND A') or text.startswith('MUSIC'):
        return 'MA'
    if text.startswith('PE') or 'PHYSICAL' in text or 'HEALTH' in text:
        return 'PEH'
    return None


def _term_values(raw, blocks, prefix, hps_row):
    """Item scores of one learner as Decimals, block by block, with the reasons any score cannot be used."""
    errors, values = [], []
    for block, items in zip(blocks, raw):
        out = []
        for i, (value, top) in enumerate(zip(items, block['hps']), start=1):
            if value in (None, ''):
                out.append(None)
                continue
            number = dec(value)
            if block['tests']:
                name = prefix + _EXAM_NAMES[min(i, 3) - 1]
            else:
                name = prefix + _TERM_PART_NAMES[block['part']] + (f' ({block["domain"]})' if block['domain'] else '') + f' {i}'
            if number is None:
                errors.append(f'{name}: "{value}" is not a number.')
            elif number < 0:
                errors.append(f'{name}: {number} is below zero.')
            elif top is None or top <= 0:
                errors.append(f'{name}: has a score but no highest possible score in row {hps_row}.')
            elif number > top:
                errors.append(f'{name}: {number} is more than the highest possible score of {top}.')
            out.append(number if number is not None and number >= 0 and top and top > 0 and number <= top else None)
        values.append(out)
    return values, errors


def _analyse_terms(parsed, period_number, *, subject_names, roster, mapeh):
    """analyse() for the three-term record: one term of one class, checked and matched (same result shape).

    MAPEH keeps Music and Arts and Physical Education and Health on two sheets a term, and its term grade is their
    rounded average. TLE kept per component has ICT and the term's other area (AFA, FCS or IA), weighed 25 and 75.
    """
    sheets = parsed['periods'].get(period_number)
    if not sheets:
        found = ', '.join(str(n) for n in sorted(parsed['periods']))
        raise EcrError(f'The file has no sheet for Term {period_number} (it has Term {found}).')
    combine = None
    if mapeh:
        parts = {term_mapeh_part(s['component'] or s['sheet']): s for s in sheets}
        if not all(p in parts for p in TERM_MAPEH_PARTS):
            raise EcrError('A MAPEH term grade needs the Music and Arts and the Physical Education and Health sheets of the '
                           'three-term MAPEH record. Download it for this class from Class records.')
        sheets, combine = [parts[p] for p in TERM_MAPEH_PARTS], 'mapeh'
    elif len(sheets) == 2 and any(_clean(s['component']) == 'ICT' for s in sheets):
        sheets, combine = sorted(sheets, key=lambda s: _clean(s['component']) != 'ICT'), 'tle'
    elif len(sheets) > 1:
        raise EcrError('The file has more than one subject for this term. Upload one subject per file.')

    checks = []
    area = learning_area(*subject_names)
    key = 'mapeh_tle' if area == 'mapeh_tle' else ('core' if area else None)
    for s in sheets:
        w, row = s['weights'], s['hps_row']
        if any(b['weight'] is None for b in s['layout']['blocks']) or abs(sum(w) - 1) > Decimal('0.0001'):
            checks.append(('error', f'{s["sheet"]}: the component weights in row {row} do not add up to 100%.'))
        elif key and w != TERM_AREA_WEIGHTS[key][1]:
            label, expected = TERM_AREA_WEIGHTS[key]
            checks.append(('error', f'{s["sheet"]} weighs written/oral works, performance tasks and examinations {_pct(w)}, '
                                    f'but {label} use {_pct(expected)} (DepEd Order No. 015, s. 2026). This looks like the '
                                    'template of another subject.'))
        for block in s['layout']['blocks']:
            if block['weight'] and not any(h for h in block['hps'] if h):
                what = 'the examinations' if block['tests'] else _TERM_PART_NAMES[block['part']].lower() + 's'
                checks.append(('warn', f'{s["sheet"]}: there is no highest possible score for {what} in row {row}.'))

    rows = []
    for base in sheets[0]['rows']:
        entry = {'row': base['row'], 'sex': base['sex'], 'name': base['name'], 'lrn': base['lrn'],
                 'errors': [], 'warnings': [], 'notes': [], 'parts': []}
        for s in sheets:
            blocks = s['layout']['blocks']
            same = next((x for x in s['rows'] if x['row'] == base['row']), None)
            prefix = f'{s["component"]}: ' if combine else ''
            if same is None:
                entry['parts'].append({'component': s['component'], 'has_scores': False,
                                       'result': compute_term(blocks, [[None] * len(b['items']) for b in blocks])})
                continue
            values, errors = _term_values(same['raw'], blocks, prefix, s['hps_row'])
            result = compute_term(blocks, values)
            entry['errors'] += errors
            if same['file_quarterly'] is not None and result['quarterly'] is not None and (
                    same['file_quarterly'] != result['quarterly'] or
                    (same['file_initial'] is not None and abs(same['file_initial'] - result['initial']) > Decimal('0.01'))):
                entry['warnings'].append(f'{prefix}the file shows {same["file_quarterly"]}, but its scores give '
                                         f'{result["quarterly"]} (initial grade {result["initial"]}). The formulas may have been changed.')
            entry['parts'].append({'component': s['component'], 'result': result,
                                   'has_scores': any(v is not None for vals in values for v in vals)})
        results = [p['result'] for p in entry['parts']]
        entry['has_scores'] = any(p['has_scores'] for p in entry['parts'])
        graded = [r for r in results if r['quarterly'] is not None]
        if combine == 'mapeh':                     # =IF(COUNT(F:G)=0,"",ROUND(AVERAGE(F:G),0))
            entry['quarterly'] = int(xround(Decimal(sum(r['quarterly'] for r in graded)) / len(graded), 0)) if graded else None
            entry['initial'] = xround(sum(r['initial'] for r in graded) / len(graded)) if graded else None
            if len(graded) == 1:
                entry['notes'].append('Only one MAPEH component has scores, so the record\'s term grade is that one alone.')
        elif combine == 'tle':                     # =ROUND((ICT*0.25)+(other*0.75),0): both are needed
            if len(graded) == 2:
                ict, other = results
                entry['quarterly'] = int(xround(ict['quarterly'] * TERM_ICT_SHARE + other['quarterly'] * (1 - TERM_ICT_SHARE), 0))
                entry['initial'] = xround(ict['initial'] * TERM_ICT_SHARE + other['initial'] * (1 - TERM_ICT_SHARE))
            else:
                entry['quarterly'] = entry['initial'] = None
                if graded:
                    entry['notes'].append('A TLE term grade needs both of the term\'s components, so there is none yet.')
        else:
            entry['initial'], entry['quarterly'] = results[0]['initial'], results[0]['quarterly']
        if entry['errors']:
            entry['initial'] = entry['quarterly'] = None          # a score cannot be trusted, so neither can the grade
        if entry['has_scores'] and not entry['name']:
            entry['errors'].append('There are scores in this row but no learner name.')
        rows.append(entry)

    named = [e for e in rows if e['name']]
    match_rows(named, roster)
    for e in rows:
        if not e['name']:
            e.update(status='unmatched', enrollment=None, candidates=[], reasons=[])
    other = sorted(n for n, ss in parsed['periods'].items() if n != period_number and
                   any(any(v not in (None, '') for vals in r['raw'] for v in vals) for s in ss for r in s['rows']))
    return {'sheets': [s['sheet'] for s in sheets], 'checks': checks, 'rows': rows, 'other_periods': other,
            'weights': sheets[0]['weights']}
