"""Report-style Excel downloads for ARMS.

A sheet gets a title block (school, report title, what it covers, when it was made), optional
key-figure boxes, short text sections and a table with a colored header row, zebra rows, number
formats, proficiency-level colors, a frozen header, filter buttons and print settings (landscape,
fit to one page wide, header row repeated on every printed page, page numbers).
"""
from datetime import datetime
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break

BRAND = '123499'
INK = '00072D'
TEXT = '2D3748'
MUTED = '4A5568'
LINE = 'D9DEE8'
ZEBRA = 'F5F7FB'
SOFT = 'EEF2FF'
FONT = 'Calibri'

# fill and text color for each proficiency level (DepEd Order No. 31, s. 2012)
LEVEL_COLORS = {
    'beginning': ('FDECEC', 'C62828'),
    'developing': ('FFF3E0', 'B45309'),
    'approaching': ('E8EAF6', '1565C0'),
    'proficient': ('E3F2FD', '123499'),
    'advanced': ('E8F5E9', '2E7D32'),
}
TONE_COLORS = {'risk': 'C62828', 'warn': 'B45309', 'good': '2E7D32', 'info': BRAND}

_thin = Side(style='thin', color=LINE)
_BOX = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
_UNDER = Border(bottom=_thin)


def _fill(color):
    return PatternFill('solid', start_color=color, end_color=color)


class Sheet:
    """One worksheet that is filled from top to bottom."""

    def __init__(self, ws, width, school_name, title, subtitle, meta):
        self.ws, self.width, self.row = ws, width, 1
        ws.sheet_view.showGridLines = False
        self._line(school_name.upper(), Font(name=FONT, size=9, bold=True, color=MUTED))
        self._line(title, Font(name=FONT, size=16, bold=True, color=BRAND), height=26)
        if subtitle:
            self._line(subtitle, Font(name=FONT, size=10.5, color=TEXT))
        for i in range(0, len(meta), 2):
            pair = meta[i:i + 2]
            self._meta(pair)
        self.row += 1
        for c in range(1, width + 1):
            ws.cell(row=self.row - 1, column=c).border = Border(bottom=Side(style='medium', color=BRAND))
        self.row += 1

    # -- helpers --------------------------------------------------------------------------
    def _line(self, text, font, height=None):
        ws = self.ws
        ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=self.width)
        cell = ws.cell(row=self.row, column=1, value=text)
        cell.font, cell.alignment = font, Alignment(vertical='center')
        if height:
            ws.row_dimensions[self.row].height = height
        self.row += 1

    def _meta(self, pairs):
        ws, half = self.ws, max(1, self.width // 2)
        for i, (label, value) in enumerate(pairs):
            start = 1 + i * half
            end = self.width if i == 1 else max(start, half)
            ws.merge_cells(start_row=self.row, start_column=start, end_row=self.row, end_column=end)
            cell = ws.cell(row=self.row, column=start)
            cell.value = f'{label}:  {value}'
            cell.font = Font(name=FONT, size=10, color=TEXT)
        self.row += 1

    # -- content --------------------------------------------------------------------------
    def page_break(self):
        """Start a new printed page at the current row."""
        self.ws.row_breaks.append(Break(id=self.row - 1))

    def heading(self, text):
        self.row += 1
        ws = self.ws
        ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=self.width)
        cell = ws.cell(row=self.row, column=1, value=text)
        cell.font = Font(name=FONT, size=12, bold=True, color=INK)
        for c in range(1, self.width + 1):
            ws.cell(row=self.row, column=c).border = _UNDER
        self.row += 1

    def figures(self, items, per_row=4):
        """Key figures as boxes: [(label, value, tone or None)], `per_row` boxes to a row."""
        ws = self.ws
        for start in range(0, len(items), per_row):
            chunk = items[start:start + per_row]
            span = max(1, self.width // len(chunk))
            for i, (label, value, tone) in enumerate(chunk):
                c1 = 1 + i * span
                c2 = self.width if i == len(chunk) - 1 else c1 + span - 1
                for r, (text, font) in enumerate((
                        (label.upper(), Font(name=FONT, size=8.5, bold=True, color=MUTED)),
                        (value, Font(name=FONT, size=14, bold=True, color=TONE_COLORS.get(tone, INK))))):
                    ws.merge_cells(start_row=self.row + r, start_column=c1, end_row=self.row + r, end_column=c2)
                    cell = ws.cell(row=self.row + r, column=c1, value=text)
                    cell.font, cell.alignment = font, Alignment(horizontal='center', vertical='center')
                    for c in range(c1, c2 + 1):
                        ws.cell(row=self.row + r, column=c).fill = _fill(SOFT)
                        ws.cell(row=self.row + r, column=c).border = Border(
                            left=_thin if c == c1 else None, right=_thin if c == c2 else None,
                            top=_thin if r == 0 else None, bottom=_thin if r == 1 else None)
            ws.row_dimensions[self.row + 1].height = 22
            self.row += 3

    def paragraph(self, text, bold=False, color=TEXT, indent=0, size=10.5):
        """Wrapped text across the sheet width; row height follows the length."""
        ws = self.ws
        ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=self.width)
        cell = ws.cell(row=self.row, column=1, value=text)
        cell.font = Font(name=FONT, size=size, bold=bold, color=color)
        cell.alignment = Alignment(wrap_text=True, vertical='top', indent=indent)
        chars_per_line = max(40, int(sum(ws.column_dimensions[get_column_letter(c)].width or 10
                                             for c in range(1, self.width + 1)) * 1.1))
        lines = max(1, -(-len(text) // chars_per_line))
        ws.row_dimensions[self.row].height = max(16, 15 * lines + 2)
        self.row += 1

    def table(self, columns, rows):
        """columns: [{'title', 'format' ('text'|'grade'|'percent'|'int'|'delta'), 'wrap'}].

        rows: lists of values. Pass a cell as (text, level_key) to color it by proficiency level.
        """
        ws = self.ws
        self.row += 1
        header = self.row
        for i, col in enumerate(columns, start=1):
            cell = ws.cell(row=header, column=i, value=col['title'])
            cell.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
            cell.fill = _fill(BRAND)
            cell.alignment = Alignment(horizontal='left' if col.get('format', 'text') == 'text' else 'center',
                                       vertical='center', wrap_text=True)
            cell.border = _BOX
        ws.row_dimensions[header].height = 30
        for r_i, values in enumerate(rows):
            r = header + 1 + r_i
            for c_i, (col, value) in enumerate(zip(columns, values), start=1):
                level = None
                if isinstance(value, tuple):
                    value, level = value
                cell = ws.cell(row=r, column=c_i, value=value)
                fmt = col.get('format', 'text')
                cell.font = Font(name=FONT, size=10, color=TEXT)
                cell.border = Border(bottom=_thin, left=_thin if c_i == 1 else None, right=_thin if c_i == len(columns) else None)
                if r_i % 2:
                    cell.fill = _fill(ZEBRA)
                if fmt == 'grade':
                    cell.number_format = '0.0'
                elif fmt == 'percent':
                    cell.number_format = '0.0"%"'
                elif fmt == 'delta':
                    cell.number_format = '+0.0;-0.0;0.0'
                elif fmt == 'int':
                    cell.number_format = '0'
                cell.alignment = Alignment(horizontal='left' if fmt == 'text' else 'center', vertical='top',
                                           wrap_text=bool(col.get('wrap')))
                if level in LEVEL_COLORS:
                    bg, fg = LEVEL_COLORS[level]
                    cell.fill = _fill(bg)
                    cell.font = Font(name=FONT, size=10, bold=True, color=fg)
                if fmt == 'delta' and isinstance(value, (int, float)):
                    cell.font = Font(name=FONT, size=10, bold=True,
                                     color='2E7D32' if value > 0 else ('C62828' if value < 0 else MUTED))
        last = header + len(rows)
        if rows:
            ws.auto_filter.ref = f'A{header}:{get_column_letter(len(columns))}{last}'
        ws.freeze_panes = ws.cell(row=header + 1, column=1)
        ws.print_title_rows = f'{header}:{header}'
        self.row = last + 1
        return header


class Report:
    """A workbook of report sheets for one download."""

    def __init__(self, title, school_name, prepared_by=''):
        self.title, self.school_name, self.prepared_by = title, school_name, prepared_by
        self.made = timezone.localtime() if timezone.is_aware(timezone.now()) else datetime.now()
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.wb.properties.title = title
        self.wb.properties.creator = 'ARMS'
        self.wb.properties.company = school_name

    def sheet(self, name, widths, subtitle='', meta=(), title=None, landscape=True):
        """A new sheet whose columns have the given widths (in characters)."""
        ws = self.wb.create_sheet(name[:31])
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top, ws.page_margins.bottom = 0.6, 0.6
        ws.oddFooter.left.text = f'ARMS · {self.school_name}'
        ws.oddFooter.center.text = 'Page &P of &N'
        ws.oddFooter.right.text = f'Generated {self.made:%b %d, %Y %I:%M %p}'
        for part in (ws.oddFooter.left, ws.oddFooter.center, ws.oddFooter.right):
            part.size, part.color = 8, MUTED
        meta = list(meta) + [('Generated', f'{self.made:%B %d, %Y, %I:%M %p}' + (f' by {self.prepared_by}' if self.prepared_by else ''))]
        return Sheet(ws, len(widths), self.school_name, title or self.title, subtitle, meta)

    def response(self, filename):
        buffer = BytesIO()
        self.wb.save(buffer)
        response = HttpResponse(buffer.getvalue(),
                                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        response['Cache-Control'] = 'no-store'
        return response


def safe_filename(*parts):
    """A plain file name from labels: letters, digits, dashes and underscores only."""
    import re
    joined = '_'.join(p for p in parts if p)
    return re.sub(r'[^A-Za-z0-9_-]+', '-', joined).strip('-_') or 'report'
